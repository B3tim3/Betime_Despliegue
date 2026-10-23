"""
Pruebas del listado de asistencia de eventos (pantalla del checklist y
exportación a Excel / Word / PDF) con los datos del aprendiz que pide
Bienestar.

Se ejecutan con:   python manage.py test eventos.test_asistencia
"""

from datetime import date, timedelta
from io import BytesIO

from django.db import connection
from django.test import TestCase
from django.test.utils import CaptureQueriesContext
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from bienestar.services import registrar_asistencia_y_horas
from eventos.models import Evento, Inscripcion
from usuarios.models import (
    Aprendiz, Grupo, Jornada, NivelFormacion, PersonalBienestar,
    ProgramaFormacion, Usuario,
)
from usuarios.services.datos_aprendiz import DatosAprendiz

CLAVE = "x"


class AsistenciaBase(TestCase):

    @classmethod
    def setUpTestData(cls):
        nivel, _ = NivelFormacion.objects.get_or_create(nombre="Tecnólogo")
        noct = Jornada.objects.create(nombre="Nocturna")
        diur = Jornada.objects.create(nombre="Diurna")
        prog = ProgramaFormacion.objects.create(
            nombre="Análisis y Desarrollo de Software", codigo="P1",
            nivel_formacion=nivel, duracion_meses=24)
        cls.g_noct = Grupo.objects.create(
            numero_grupo="3000002", programa_formacion=prog, jornada=noct,
            fecha_inicio=date(2026, 1, 1), fecha_fin=date(2027, 1, 1))
        cls.g_diur = Grupo.objects.create(
            numero_grupo="3000001", programa_formacion=prog, jornada=diur,
            fecha_inicio=date(2026, 1, 1), fecha_fin=date(2027, 1, 1))

        staff = Usuario.objects.create_user(
            numero_documento="900", password=CLAVE, tipo_documento="CC",
            nombres="Paola", apellidos="Ramírez", email="p@x.co")
        cls.bienestar = PersonalBienestar.objects.create(
            usuario=staff, cargo="Coordinadora", fecha_ingreso=date(2026, 1, 1))
        cls.staff = staff

        def aprendiz(doc, nombres, apellidos, grupo, **extra):
            u = Usuario.objects.create_user(
                numero_documento=doc, password=CLAVE, tipo_documento=extra.pop("tipo", "CC"),
                nombres=nombres, apellidos=apellidos, **extra)
            return Aprendiz.objects.create(usuario=u, grupo=grupo, fecha_ingreso=date(2026, 1, 2))

        cls.ana = aprendiz("101", "Ana", "Gómez", cls.g_noct, tipo="TI",
                           email="ana@x.co", genero="F", telefono="3001112233")
        cls.luis = aprendiz("102", "Luis", "Vargas", cls.g_diur,
                            email="luis@x.co", genero="M", telefono="3004445566")
        # Cuenta "antigua": sin género, correo ni teléfono -> no debe romper nada.
        cls.vieja = aprendiz("103", "Zoila", "Zapata", cls.g_diur)

        ahora = timezone.now()
        cls.evento = Evento.objects.create(
            titulo="Taller: pausas / activas [2026]", lugar="Auditorio", horas_otorgadas=2,
            creado_por=staff, fecha_inicio=ahora + timedelta(days=1),
            fecha_fin=ahora + timedelta(days=1, hours=2))
        cls.ins = {a.usuario.numero_documento: Inscripcion.objects.create(evento=cls.evento, aprendiz=a)
                   for a in (cls.ana, cls.luis, cls.vieja)}

    def setUp(self):
        self.client.force_login(self.staff)

    def exportar(self, formato="excel", **extra):
        return self.client.get(reverse("eventos:panel_inscritos_exportar"),
                               {"formato": formato, "evento_id": self.evento.id, **extra})

    def filas_excel(self, respuesta):
        libro = load_workbook(BytesIO(respuesta.content))
        return [[c.value for c in fila] for fila in libro.active.iter_rows()]


class DatosAprendizTests(AsistenciaBase):

    def test_todos_los_datos_en_el_orden_de_los_encabezados(self):
        fila = DatosAprendiz.de(Aprendiz.objects.select_related(
            "usuario", "grupo__jornada", "grupo__programa_formacion__nivel_formacion").get(pk=self.ana.pk)).como_fila()
        self.assertEqual(len(fila), len(DatosAprendiz.ENCABEZADOS))
        self.assertEqual(fila, [
            "Ana Gómez", "TI", "101", "Femenino", "3000002",
            "Análisis y Desarrollo de Software", "Nocturna", "Tecnólogo", "ana@x.co", "3001112233"])

    def test_cuenta_antigua_sin_datos_nuevos_no_falla(self):
        fila = DatosAprendiz.de(self.vieja).como_fila()
        self.assertEqual((fila[3], fila[8], fila[9]), ("", "", ""))


class ChecklistPantallaTests(AsistenciaBase):

    def test_muestra_el_grupo_y_todos_los_datos(self):
        r = self.client.get(reverse("eventos:panel_inscritos", args=[self.evento.id]))
        self.assertEqual(r.status_code, 200)
        for texto in ("3000002", "3000001", "Nocturna", "Diurna", "Tecnólogo", "Femenino", "Masculino",
                      "ana@x.co", "3001112233", "Análisis y Desarrollo de Software", "TI 101"):
            self.assertContains(r, texto)

    def test_cuenta_antigua_muestra_guion(self):
        r = self.client.get(reverse("eventos:panel_inscritos", args=[self.evento.id]))
        self.assertContains(r, "dato-vacio")

    def test_boton_exportar_apunta_a_este_evento(self):
        r = self.client.get(reverse("eventos:panel_inscritos", args=[self.evento.id]))
        self.assertContains(r, f"evento_id={self.evento.id}")

    def test_pocas_consultas_a_la_base(self):
        # Sin N+1: con 3 aprendices o con 13, la cantidad de consultas es la misma.
        url = reverse("eventos:panel_inscritos", args=[self.evento.id])
        with CaptureQueriesContext(connection) as pocas:
            self.client.get(url)

        grupo = self.g_diur
        for n in range(10):
            u = Usuario.objects.create_user(
                numero_documento=f"2{n:02d}", password=CLAVE, tipo_documento="CC",
                nombres=f"Extra{n}", apellidos="Prueba")
            Inscripcion.objects.create(
                evento=self.evento,
                aprendiz=Aprendiz.objects.create(usuario=u, grupo=grupo, fecha_ingreso=date(2026, 1, 2)))

        with CaptureQueriesContext(connection) as muchas:
            self.client.get(url)

        self.assertEqual(len(pocas), len(muchas))

    def test_horas_guardadas_se_conservan_al_recargar(self):
        registrar_asistencia_y_horas(
            inscripcion=self.ins["101"], asistio=True, cantidad_horas=5,
            motivo="x", asignado_por=self.bienestar)
        r = self.client.get(reverse("eventos:panel_inscritos", args=[self.evento.id]))
        self.assertContains(r, 'value="5"')   # no vuelve al valor por defecto (2)


class ExportacionAsistenciaTests(AsistenciaBase):

    def test_excel_trae_todas_las_columnas_y_filas(self):
        r = self.exportar("excel")
        self.assertEqual(r.status_code, 200)
        filas = self.filas_excel(r)
        # Las primeras filas son el membrete; el encabezado es la que empieza por "N°".
        i = next(n for n, f in enumerate(filas) if f and f[0] == "N°")
        self.assertEqual(filas[i], ["N°", *DatosAprendiz.ENCABEZADOS, "Asistió"])
        datos = [f for f in filas[i + 1:] if f and f[0] and f[0].isdigit()]
        self.assertEqual([f[1] for f in datos], ["Ana Gómez", "Luis Vargas", "Zoila Zapata"])  # por apellido
        self.assertEqual(datos[0][2:5], ["TI", "101", "Femenino"])
        self.assertEqual(datos[0][5:9], ["3000002", "Análisis y Desarrollo de Software", "Nocturna", "Tecnólogo"])
        self.assertEqual(datos[0][-1], "Pendiente")

    def test_asistencia_real_en_el_archivo(self):
        registrar_asistencia_y_horas(
            inscripcion=self.ins["102"], asistio=True, cantidad_horas=3,
            motivo="x", asignado_por=self.bienestar)
        filas = self.filas_excel(self.exportar("excel"))
        luis = next(f for f in filas if f and f[1] == "Luis Vargas")
        self.assertEqual(luis[-1], "Sí")
        self.assertNotIn("Horas", [c for f in filas for c in f])

    def test_filtro_q_exporta_solo_lo_filtrado(self):
        filas = self.filas_excel(self.exportar("excel", q="nocturna"))
        nombres = [f[1] for f in filas if f and f[0] and f[0].isdigit()]
        self.assertEqual(nombres, ["Ana Gómez"])

    def test_nombre_del_evento_con_caracteres_raros_no_rompe_excel(self):
        # "Taller: pausas / activas [2026]" tiene : / [ ] que Excel no admite en el nombre de hoja.
        self.assertEqual(self.exportar("excel").status_code, 200)

    def test_word_y_pdf(self):
        self.assertEqual(self.exportar("word").status_code, 200)
        r = self.exportar("pdf")
        self.assertEqual(r.status_code, 200)
        self.assertTrue(r.content.startswith(b"%PDF"))

    def test_inscripcion_cancelada_no_sale(self):
        self.ins["103"].estado = "cancelada"
        self.ins["103"].save()
        nombres = [f[1] for f in self.filas_excel(self.exportar("excel")) if f and f[0] and f[0].isdigit()]
        self.assertNotIn("Zoila Zapata", nombres)

    def test_sin_evento_redirige_con_aviso(self):
        r = self.client.get(reverse("eventos:panel_inscritos_exportar"), {"formato": "excel"})
        self.assertRedirects(r, reverse("eventos:panel_lista"), fetch_redirect_response=False)

    def test_un_aprendiz_no_puede_exportar(self):
        self.client.force_login(self.ana.usuario)
        r = self.exportar("excel")
        self.assertEqual(r.status_code, 302)


class ExportacionAprendicesTests(AsistenciaBase):

    def test_listado_de_aprendices_trae_los_mismos_datos(self):
        r = self.client.get(reverse("usuarios:aprendices_exportar"), {"formato": "excel"})
        self.assertEqual(r.status_code, 200)
        filas = self.filas_excel(r)
        i = next(n for n, f in enumerate(filas) if f and f[0] == "Nombres y apellidos")
        self.assertEqual(filas[i], [*DatosAprendiz.ENCABEZADOS, "Estado"])
        ana = next(f for f in filas[i + 1:] if f and f[0] == "Ana Gómez")
        self.assertEqual(ana[3:8], ["Femenino", "3000002", "Análisis y Desarrollo de Software", "Nocturna", "Tecnólogo"])
        self.assertEqual(ana[-1], "Activo")
