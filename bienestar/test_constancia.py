"""
Pruebas del historial de asistencia y de la constancia de horas en PDF
(emisión, descarga, verificación pública del código).

Se ejecutan con:   python manage.py test bienestar.test_constancia
"""

from datetime import date, timedelta
from io import BytesIO

from django.db import connection
from django.test import TestCase
from django.test.utils import CaptureQueriesContext
from django.urls import reverse
from django.utils import timezone

from bienestar.models import ConfiguracionBienestar, ConstanciaHoras
from bienestar.services import registrar_asistencia_y_horas
from bienestar.services.constancia import SinHorasError, emitir_constancia, generar_pdf_constancia
from bienestar.services.historial import historial_de, resumen_historial
from bienestar.services.progreso import horas_de, total_horas_de
from eventos.models import Evento, Inscripcion
from usuarios.models import (
    Administrador, Grupo, Jornada, NivelFormacion, PersonalBienestar, ProgramaFormacion, Usuario,
)
from usuarios.services import registrar_aprendiz


class Base(TestCase):

    @classmethod
    def setUpTestData(cls):
        jornada = Jornada.objects.create(nombre="Diurna")
        nivel, _ = NivelFormacion.objects.get_or_create(nombre="Tecnólogo")
        programa = ProgramaFormacion.objects.create(
            nombre="Análisis y Desarrollo de Software", codigo="228118",
            nivel_formacion=nivel, duracion_meses=24)
        cls.grupo = Grupo.objects.create(
            numero_grupo="3171693", programa_formacion=programa, jornada=jornada,
            fecha_inicio=date(2026, 1, 1), fecha_fin=date(2027, 12, 31))

        cls.usuario, cls.aprendiz = registrar_aprendiz(
            tipo_documento="CC", numero_documento="1012345678", nombres="Juan",
            apellidos="Rodríguez", email="juan@example.com", password="ClaveSegura#2026",
            grupo=cls.grupo, genero="M", telefono="3001234567")
        cls.otro_usuario, cls.otro = registrar_aprendiz(
            tipo_documento="CC", numero_documento="1099999999", nombres="Otra",
            apellidos="Persona", email="otra@example.com", password="ClaveSegura#2026",
            grupo=cls.grupo, genero="F", telefono="3109999999")

        cls.staff_user = Usuario.objects.create_user(
            numero_documento="900000002", password="ClaveSegura#2026", tipo_documento="CC",
            nombres="Paola", apellidos="Bienestar", email="paola@example.com")
        cls.personal = PersonalBienestar.objects.create(
            usuario=cls.staff_user, cargo="Profesional", fecha_ingreso=date(2026, 1, 1))

        cls.admin = Usuario.objects.create_user(
            numero_documento="900000001", password="ClaveSegura#2026", tipo_documento="CC",
            nombres="Admin", apellidos="Be Time", email="admin@example.com")
        Administrador.objects.create(usuario=cls.admin, fecha_ingreso=date(2026, 1, 1))

    def evento(self, titulo="Taller", dias=0, horas=4, estado="finalizado"):
        inicio = timezone.now() + timedelta(days=dias)
        return Evento.objects.create(
            titulo=titulo, descripcion="x", lugar="Aula", fecha_inicio=inicio,
            fecha_fin=inicio + timedelta(hours=2), horas_otorgadas=horas,
            creado_por=self.staff_user, estado=estado)

    def inscribir(self, evento, aprendiz=None):
        return Inscripcion.objects.create(evento=evento, aprendiz=aprendiz or self.aprendiz)

    def dar_horas(self, inscripcion, horas=4, asistio=True):
        registrar_asistencia_y_horas(
            inscripcion=inscripcion, asistio=asistio, cantidad_horas=horas,
            motivo="x", asignado_por=self.personal)


class ReglaDeHorasTests(Base):

    def test_solo_cuentan_las_horas_de_asistencia_marcada(self):
        a = self.inscribir(self.evento("A", -3))
        b = self.inscribir(self.evento("B", -2))
        self.dar_horas(a, 4)
        self.dar_horas(b, 6)
        self.assertEqual(total_horas_de(self.aprendiz), 10)
        # Bienestar corrige: B no asistió. Las horas ya no cuentan (sin borrar el registro).
        self.dar_horas(b, 6, asistio=False)
        self.assertEqual(total_horas_de(self.aprendiz), 4)
        self.assertEqual(horas_de(self.aprendiz).count(), 1)

    def test_no_mezcla_aprendices(self):
        self.dar_horas(self.inscribir(self.evento("A", -3)), 4)
        self.assertEqual(total_horas_de(self.otro), 0)

    def test_dashboard_usa_la_misma_regla(self):
        a = self.inscribir(self.evento("A", -3))
        self.dar_horas(a, 4)
        self.dar_horas(a, 4, asistio=False)
        self.client.force_login(self.usuario)
        r = self.client.get(reverse("usuarios:dashboard_aprendiz"))
        self.assertEqual(r.context["progreso"]["total"], 0)


class EmisionTests(Base):

    def test_sin_horas_no_se_emite(self):
        with self.assertRaises(SinHorasError):
            emitir_constancia(self.aprendiz)
        self.assertEqual(ConstanciaHoras.objects.count(), 0)

    def test_guarda_la_foto_de_lo_certificado(self):
        self.dar_horas(self.inscribir(self.evento("A", -3)), 12)
        c = emitir_constancia(self.aprendiz)
        self.assertEqual((c.nombre, c.numero_documento, c.programa, c.grupo, c.total_horas, c.meta_horas),
                         ("Juan Rodríguez", "1012345678", "Análisis y Desarrollo de Software", "3171693", 12, 40))
        self.assertRegex(c.codigo, r"^BT-[A-Z2-9]{4}-[A-Z2-9]{4}$")

    def test_la_constancia_vieja_no_cambia_aunque_cambien_los_datos(self):
        self.dar_horas(self.inscribir(self.evento("A", -3)), 12)
        c = emitir_constancia(self.aprendiz)
        self.usuario.nombres = "Cambiado"
        self.usuario.save()
        ConfiguracionBienestar(meta_horas=100).save()
        c.refresh_from_db()
        self.assertEqual((c.nombre, c.meta_horas), ("Juan Rodríguez", 40))

    def test_mismo_dia_y_mismas_horas_reutiliza_el_codigo(self):
        self.dar_horas(self.inscribir(self.evento("A", -3)), 12)
        c1 = emitir_constancia(self.aprendiz)
        c2 = emitir_constancia(self.aprendiz)
        self.assertEqual(c1.pk, c2.pk)
        self.assertEqual(ConstanciaHoras.objects.count(), 1)

    def test_si_cambian_las_horas_se_emite_una_nueva(self):
        self.dar_horas(self.inscribir(self.evento("A", -3)), 12)
        c1 = emitir_constancia(self.aprendiz)
        self.dar_horas(self.inscribir(self.evento("B", -2)), 4)
        c2 = emitir_constancia(self.aprendiz)
        self.assertNotEqual(c1.codigo, c2.codigo)
        self.assertEqual(c2.total_horas, 16)

    def test_documento_oculto(self):
        self.dar_horas(self.inscribir(self.evento("A", -3)), 12)
        self.assertEqual(emitir_constancia(self.aprendiz).documento_oculto, "••••••5678")


class DescargaPdfTests(Base):

    def setUp(self):
        self.url = reverse("bienestar:constancia")

    def test_aprendiz_con_horas_descarga_el_pdf(self):
        self.dar_horas(self.inscribir(self.evento("A", -3)), 12)
        self.client.force_login(self.usuario)
        r = self.client.get(self.url)
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r["Content-Type"], "application/pdf")
        self.assertTrue(r.content.startswith(b"%PDF"))
        c = ConstanciaHoras.objects.get()
        self.assertIn(c.codigo, r["Content-Disposition"])

    def test_titulo_de_evento_con_caracteres_raros_no_rompe_el_pdf(self):
        self.dar_horas(self.inscribir(self.evento("Taller <b>& \"pausas\" > 3", -3)), 12)
        self.client.force_login(self.usuario)
        r = self.client.get(self.url)
        self.assertTrue(r.content.startswith(b"%PDF"))

    def test_muchos_eventos_ocupan_varias_paginas_sin_error(self):
        for i in range(60):
            self.dar_horas(self.inscribir(self.evento(f"Evento {i}", -i - 1)), 1)
        c = emitir_constancia(self.aprendiz)
        pdf = generar_pdf_constancia(c, "http://testserver/x/")
        self.assertTrue(pdf.startswith(b"%PDF"))
        self.assertEqual(c.total_horas, 60)

    def test_sin_horas_redirige_con_aviso(self):
        self.client.force_login(self.usuario)
        r = self.client.get(self.url)
        self.assertRedirects(r, reverse("usuarios:dashboard"), fetch_redirect_response=False)
        self.assertEqual(ConstanciaHoras.objects.count(), 0)

    def test_quien_no_es_aprendiz_es_redirigido(self):
        self.client.force_login(self.staff_user)
        r = self.client.get(self.url)
        self.assertRedirects(r, reverse("usuarios:dashboard"), fetch_redirect_response=False)

    def test_requiere_sesion(self):
        r = self.client.get(self.url)
        self.assertEqual(r.status_code, 302)
        self.assertIn(reverse("usuarios:login"), r["Location"])

    def test_la_constancia_siempre_es_la_del_usuario_de_la_sesion(self):
        self.dar_horas(self.inscribir(self.evento("A", -3), self.otro), 8)
        self.client.force_login(self.usuario)      # Juan no tiene horas; la otra sí
        r = self.client.get(self.url)
        self.assertEqual(r.status_code, 302)
        self.assertEqual(ConstanciaHoras.objects.count(), 0)


class VerificacionTests(Base):

    def setUp(self):
        self.dar_horas(self.inscribir(self.evento("A", -3)), 12)
        self.c = emitir_constancia(self.aprendiz)

    def url(self, codigo):
        return reverse("bienestar:verificar_constancia", args=[codigo])

    def test_codigo_valido_es_publico_y_oculta_el_documento(self):
        r = self.client.get(self.url(self.c.codigo))     # sin sesión
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, "Constancia auténtica")
        self.assertContains(r, "Juan Rodríguez")
        self.assertContains(r, "12 de 40")
        self.assertContains(r, "••••••5678")
        self.assertNotContains(r, "1012345678")
        self.assertNotContains(r, "juan@example.com")

    def test_acepta_minusculas_y_espacios(self):
        r = self.client.get(self.url(self.c.codigo.lower()))
        self.assertContains(r, "Constancia auténtica")

    def test_codigo_inexistente(self):
        r = self.client.get(self.url("BT-AAAA-AAAA"))
        self.assertEqual(r.status_code, 404)
        self.assertContains(r, "No encontramos esa constancia", status_code=404)

    def test_codigo_con_html_se_escapa(self):
        r = self.client.get("/bienestar/constancia/verificar/%3Cscript%3Ex%3C/script%3E/")
        self.assertNotContains(r, "<script>x</script>", status_code=404)


class HistorialTests(Base):

    def setUp(self):
        self.client.force_login(self.usuario)
        self.url = reverse("bienestar:mi_historial")

    def armar(self):
        self.i_asistio = self.inscribir(self.evento("Asistido", -5))
        self.dar_horas(self.i_asistio, 4)
        self.i_no = self.inscribir(self.evento("Faltado", -4))
        self.dar_horas(self.i_no, 0, asistio=False)
        self.i_pend = self.inscribir(self.evento("Pendiente", 3, estado="programado"))
        self.i_canc = self.inscribir(self.evento("Cancelada", 4, estado="programado"))
        Inscripcion.objects.filter(pk=self.i_canc.pk).update(estado="cancelada")
        self.inscribir(self.evento("Evento cancelado", 5, estado="cancelado"))
        self.inscribir(self.evento("De otra", -2), self.otro)

    def test_filas_estados_y_orden(self):
        self.armar()
        filas = historial_de(self.aprendiz)
        self.assertEqual([f.evento for f in filas],
                         ["Evento cancelado", "Cancelada", "Pendiente", "Faltado", "Asistido"])
        estados = {f.evento: f.estado_codigo for f in filas}
        self.assertEqual(estados, {
            "Evento cancelado": "evento_cancelado", "Cancelada": "cancelada",
            "Pendiente": "inscrito", "Faltado": "no_asistio", "Asistido": "asistio"})

    def test_horas_solo_si_asistio(self):
        self.armar()
        horas = {f.evento: f.horas for f in historial_de(self.aprendiz)}
        self.assertEqual(horas["Asistido"], 4)
        self.assertEqual(horas["Faltado"], 0)

    def test_resumen_no_cuenta_canceladas(self):
        self.armar()
        r = resumen_historial(historial_de(self.aprendiz))
        self.assertEqual(r, {"inscritos": 3, "asistidos": 1, "no_asistidos": 1, "pendientes": 1, "horas": 4})

    def test_pagina(self):
        self.armar()
        r = self.client.get(self.url)
        self.assertEqual(r.status_code, 200)
        for t in ("Asistido", "Faltado", "Pendiente", "Evento cancelado", "Constancia"):
            self.assertContains(r, t)
        self.assertNotContains(r, "De otra")          # nunca eventos de otro aprendiz
        self.assertEqual(r.context["resumen"]["horas"], 4)

    def test_pagina_vacia(self):
        r = self.client.get(self.url)
        self.assertContains(r, "Todavía no te has inscrito")
        self.assertNotContains(r, reverse("bienestar:constancia"))   # sin horas no hay botón

    def test_consultas_no_crecen_con_los_eventos(self):
        for i in range(3):
            self.dar_horas(self.inscribir(self.evento(f"E{i}", -i - 1)), 1)
        self.client.get(self.url)    # calienta (crea la configuración la primera vez)
        with CaptureQueriesContext(connection) as pocas:
            self.client.get(self.url)
        for i in range(3, 13):
            self.dar_horas(self.inscribir(self.evento(f"E{i}", -i - 1)), 1)
        with CaptureQueriesContext(connection) as muchas:
            self.client.get(self.url)
        self.assertEqual(len(pocas), len(muchas))

    def test_requiere_ser_aprendiz(self):
        self.client.force_login(self.staff_user)
        r = self.client.get(self.url)
        self.assertRedirects(r, reverse("usuarios:dashboard"), fetch_redirect_response=False)
        self.client.logout()
        self.assertEqual(self.client.get(self.url).status_code, 302)

    def test_enlaces_en_dashboard_y_menu(self):
        self.dar_horas(self.inscribir(self.evento("A", -3)), 4)
        r = self.client.get(reverse("usuarios:dashboard_aprendiz"))
        self.assertContains(r, self.url)
        self.assertContains(r, reverse("bienestar:constancia"))


class ExportarHistorialTests(Base):

    def setUp(self):
        self.client.force_login(self.usuario)
        self.url = reverse("bienestar:mi_historial_exportar")
        self.dar_horas(self.inscribir(self.evento("Asistido", -5)), 4)
        self.dar_horas(self.inscribir(self.evento("Faltado", -4)), 0, asistio=False)
        self.inscribir(self.evento("De otra", -2), self.otro)

    def filas_excel(self, **params):
        from openpyxl import load_workbook
        r = self.client.get(self.url, {"formato": "excel", **params})
        self.assertEqual(r.status_code, 200)
        hoja = load_workbook(BytesIO(r.content)).active
        return [[c.value for c in fila] for fila in hoja.iter_rows()]

    def test_excel_con_las_columnas_y_solo_mis_eventos(self):
        filas = self.filas_excel()
        i = next(n for n, f in enumerate(filas) if f and f[0] == "Evento")
        self.assertEqual(filas[i], ["Evento", "Lugar", "Fecha", "Estado", "Horas"])
        eventos = [f[0] for f in filas[i + 1:] if f and f[0]]
        self.assertCountEqual(eventos, ["Asistido", "Faltado"])

    def test_filtros(self):
        filas = self.filas_excel(estado="asistio")
        self.assertEqual([f[0] for f in filas if f and f[0] in ("Asistido", "Faltado")], ["Asistido"])
        filas = self.filas_excel(q="falt")
        self.assertEqual([f[0] for f in filas if f and f[0] in ("Asistido", "Faltado")], ["Faltado"])

    def test_word_y_pdf(self):
        for formato in ("word", "pdf"):
            self.assertEqual(self.client.get(self.url, {"formato": formato}).status_code, 200)

    def test_requiere_ser_aprendiz(self):
        self.client.force_login(self.staff_user)
        r = self.client.get(self.url, {"formato": "excel"})
        self.assertEqual(r.status_code, 302)
